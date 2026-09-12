"""
Generate the layer metainfo documents from metadata.xlsx.

Two outputs per layer, from one pass over the spreadsheet:

- `resultaten/<id>.html` — a standalone page for checking a document in a
  browser without running the app.
- `fragmenten/<id>.html` — what the app fetches and injects into its metainfo
  dialog. Copied into `public/data/meta/` by the same run.

Run it with `npm run metainfo`. See README.md for the one-time dependency
install.
"""

import re
import shutil
import sys
import warnings
from pathlib import Path

try:
    import openpyxl
except ModuleNotFoundError:
    sys.exit(
        "openpyxl is required to read metadata.xlsx.\n"
        "Install it with:  pip install openpyxl pillow"
    )

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

BASE_DIR = Path(__file__).resolve().parent
REPO_ROOT = BASE_DIR.parent.parent

EXCEL_PATH = BASE_DIR / "metadata.xlsx"
TEMPLATE_PATH = BASE_DIR / "template" / "opmaak.html"
FRAGMENT_TEMPLATE_PATH = BASE_DIR / "template" / "fragment.html"
OUTPUT_DIR = BASE_DIR / "resultaten"
FRAGMENT_DIR = BASE_DIR / "fragmenten"
AFBEELDINGEN_DIR = BASE_DIR / "afbeeldingen"

# Where the app serves the fragments and their images from.
PUBLIC_META_DIR = REPO_ROOT / "public" / "data" / "meta"
PUBLIC_IMG_DIR = PUBLIC_META_DIR / "afbeeldingen"

# Image paths inside a fragment must be app-absolute: a relative path would
# resolve against whatever URL the single-page app happens to be showing.
FRAGMENT_IMG_BASE = "/data/meta/afbeeldingen/"
# ...while the preview pages sit next to the images and stay relative.
PREVIEW_IMG_BASE = "../afbeeldingen/"

# The explanatory diagrams are ~1300px PNGs of a few MB each, shown a few
# hundred pixels wide in a dialog. Downscaling and converting to WebP cuts them
# by about 95% with nothing visible lost at that size.
DIAGRAM_MAX_WIDTH = 800
DIAGRAM_QUALITY = 82

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
FRAGMENT_DIR.mkdir(parents=True, exist_ok=True)


def lees_sheet_data(wb, sheet_name):
    real_sheet_name = next(
        (name for name in wb.sheetnames
         if name.strip().lower() == sheet_name.lower()),
        None,
    )
    if not real_sheet_name:
        return []

    rows = list(wb[real_sheet_name].iter_rows(values_only=True))
    if not rows:
        return []

    header_idx = None
    headers = []
    for idx, row in enumerate(rows[:10]):
        cleaned = [str(c).strip() if c is not None else "" for c in row]
        if sum(1 for c in cleaned if c) >= 2:
            headers = cleaned
            header_idx = idx
            break

    if header_idx is None:
        return []

    data = []
    for row in rows[header_idx + 1:]:
        row_dict = {}
        for h, val in zip(headers, row):
            if h:
                str_val = str(val).strip() if val is not None else ""
                row_dict[h] = "" if str_val.lower() in ("nan", "none") else str_val
        if any(row_dict.values()):
            data.append(row_dict)

    return data


def escape_html(value):
    value = str(value or "")
    return (value.replace("&", "&amp;")
                 .replace("<", "&lt;")
                 .replace(">", "&gt;")
                 .replace('"', "&quot;"))


def format_text(value):
    """
    Metadata descriptions may contain HTML such as <a> and <br>.
    All links receive the standard link styling automatically.
    """
    html = str(value or "")

    def style_link(match):
        tag = match.group(0)

        # Als de link al een class heeft, voeg de standaard classes toe
        if re.search(r'\bclass\s*=', tag, re.IGNORECASE):
            tag = re.sub(
                r'class\s*=\s*"([^"]*)"',
                lambda m: f'class="{m.group(1)} underline hover:text-blue-800"',
                tag,
                count=1,
                flags=re.IGNORECASE,
            )
        else:
            tag = tag.replace(
                "<a",
                '<a class="underline hover:text-blue-800"',
                1,
            )

        # Als er nog geen style staat, voeg de primaire kleur toe
        if not re.search(r'\bstyle\s*=', tag, re.IGNORECASE):
            tag = tag.replace(
                ">",
                ' style="color: var(--primary-color);">',
                1,
            )

        return tag

    html = re.sub(
        r"<a\b[^>]*>",
        style_link,
        html,
        flags=re.IGNORECASE,
    )

    return html


def bron_logo_url(logo, basis):
    """
    A source logo's path, rooted at `basis`.

    The spreadsheet's Logo column already carries a "../afbeeldingen/" prefix,
    written for the preview pages. Strip it before re-prefixing, or the result
    is "../afbeeldingen/../afbeeldingen/iconen/osm.png" — which happens to
    resolve, but only because the two halves cancel out.
    """
    if not logo:
        return ""
    logo = str(logo).strip()

    # Absolute URL: use unchanged.
    if re.match(r"^https?://", logo, re.I):
        return logo

    logo = re.sub(r"^(\.\./)*afbeeldingen/", "", logo.lstrip("/\\"), flags=re.I)
    return basis + logo


def herschrijf_afbeeldingen(html, basis):
    """
    Point the spreadsheet's inline <img> sources at `basis`.

    The kaartlaagtype descriptions embed their diagrams as HTML written for the
    preview pages, so the paths are relative to `resultaten/`. A fragment is
    served from somewhere else entirely and needs its own root. The diagrams are
    also converted to WebP on the way out, so the extension moves with them.
    """
    html = re.sub(
        r'(src\s*=\s*")(?:\.\./)*afbeeldingen/',
        lambda m: m.group(1) + basis,
        html,
        flags=re.I,
    )
    return re.sub(
        rf'(src\s*=\s*"{re.escape(basis)}kaartlaagtypen/[^"]+)\.png"',
        r'\1.webp"',
        html,
        flags=re.I,
    )


def maak_bronnen_html(kaartlaag_id, df_kaartlagen_bronnen, bronnen_dict, basis):
    gekoppelde_bronnen = [
        kb for kb in df_kaartlagen_bronnen
        if kb.get("Naam van de laag", "").strip() == kaartlaag_id
    ]

    if not gekoppelde_bronnen:
        return """
        <div class="rounded-xl border border-gray-200 bg-gray-50 p-4 text-sm text-gray-500">
          Voor deze kaartlaag zijn geen databronnen geregistreerd.
        </div>
        """

    html = []

    for kb in gekoppelde_bronnen:
        b_naam = (kb.get("Naam van de bron", "") or "").strip()
        datum = kb.get("Datum", "")
        b_info = bronnen_dict.get(b_naam, {})

        houder = b_info.get("Bronhouder", "")
        beschrijving = b_info.get("Beschrijving", "")
        url = b_info.get("URL", "")
        logo = bron_logo_url(b_info.get("Logo", ""), basis)

        titel = escape_html(b_naam)
        houder_html = escape_html(houder)
        datum_html = escape_html(datum)
        link_html = escape_html(url)

        if link_html:
            titel_html = f"""
                  <div class="flex items-center justify-between gap-2 mb-0.5">
                    <h4 class="font-semibold text-gray-900 text-sm">{titel}</h4>
                    <a href="{link_html}" target="_blank" rel="noopener"
                       class="text-xs font-semibold hover:underline flex items-center gap-1"
                       style="color: var(--primary-color);">
                      Bekijk bron &rsaquo;
                    </a>
                  </div>
            """
        else:
            titel_html = f"""
                  <h4 class="font-semibold text-gray-900 text-sm mb-0.5">{titel}</h4>
            """

        logo_html = ""
        if logo:
            logo_html = f"""
                    <img src="{escape_html(logo)}"
                         alt="Logo {titel}"
                         class="max-h-full max-w-full object-contain"
                         onerror="this.style.display='none';">
            """

        # No `even:bg-gray-100` here. Tailwind only scans source files, and these
        # documents are fetched at runtime, so a utility used nowhere else in the
        # app is purged from the build and the striping silently disappears. The
        # app styles these rows itself, scoped to `.dro-meta` in src/index.css.
        html.append(f"""
			<tr class="border-b border-gray-200">
			  <td class="align-top pr-4 py-3">
				{logo_html}
			  </td>

			  <td class="align-top p-3">
				{titel_html}

				<div class="text-xs text-gray-500 mb-1">
				  <span class="font-medium text-gray-700">Bronhouder:</span> {houder_html}
				  <span aria-hidden="true"> | </span>
				  <span class="font-medium text-gray-700">Datum:</span> {datum_html}
				</div>

				<p class="text-xs leading-relaxed text-gray-600">
				  {format_text(beschrijving)}
				</p>
			  </td>
			</tr>        """)

    return "\n".join(html)


def vervang_template(html, key, content):
    pattern_if = re.compile(
        rf"\{{%\s*if\s+{re.escape(key)}\s*%\}}(.*?)\{{%\s*endif\s*%\}}",
        re.DOTALL,
    )

    if content and str(content).strip():
        html = pattern_if.sub(r"\1", html)
        html = re.sub(
            rf"\{{\{{\s*{re.escape(key)}\s*\}}\}}",
            lambda m: str(content),
            html,
        )
    else:
        html = pattern_if.sub("", html)
        html = re.sub(
            rf"\{{\{{\s*{re.escape(key)}\s*\}}\}}",
            "",
            html,
        )

    return html


# ------------------------------------------------------------
# INLEZEN
# ------------------------------------------------------------

if not EXCEL_PATH.exists():
    raise FileNotFoundError(f"Excel-bestand niet gevonden: {EXCEL_PATH}")

if not TEMPLATE_PATH.exists():
    raise FileNotFoundError(f"Template-bestand niet gevonden: {TEMPLATE_PATH}")

if not FRAGMENT_TEMPLATE_PATH.exists():
    raise FileNotFoundError(f"Template-bestand niet gevonden: {FRAGMENT_TEMPLATE_PATH}")

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

df_kaartlaag = lees_sheet_data(wb, "kaartlaag")
df_kaartlaagtypen = lees_sheet_data(wb, "kaartlaagtypen")
df_bronnen = lees_sheet_data(wb, "bronnen")
df_kaartlagen_bronnen = lees_sheet_data(wb, "kaartlagen_bronnen")

kaartlaagtypen_dict = {
    row.get("Naam van het kaartlaagtype", ""): row.get("Beschrijving", "")
    for row in df_kaartlaagtypen
    if row.get("Naam van het kaartlaagtype")
}

bronnen_dict = {
    (row.get("Naam de bron", "") or row.get("Naam van de bron", "")).strip(): row
    for row in df_bronnen
    if (row.get("Naam de bron", "") or row.get("Naam van de bron", ""))
}

template_str = TEMPLATE_PATH.read_text(encoding="utf-8")
fragment_template_str = FRAGMENT_TEMPLATE_PATH.read_text(encoding="utf-8")

count = 0
skipped = 0

for row in df_kaartlaag:
    layer_id = row.get("Naam van de laag", "").strip()

    if (
        not layer_id
        or layer_id.lower().startswith("algemene informatie")
        or layer_id.lower().startswith("naam van")
    ):
        continue

    safe_filename = re.sub(r'[\\/*?:"<>|]', "_", layer_id).strip()

    titel = row.get("Titel") or layer_id
    thema = (row.get("Thema / Subthema") or "Algemeen").replace(";", " ›")
    wat_ziet_u = row.get("Wat ziet u?", "")
    kaartlaagtype_id = row.get("Gekozen kaartlaagtype", "")
    berekenwijze = kaartlaagtypen_dict.get(kaartlaagtype_id, "")
    aannames = row.get("Aannames en Onzekerheden", "")

    # A layer with no "Wat ziet u?" has only a title and a theme in the
    # spreadsheet. A dialog showing nothing but those reads as broken, whereas
    # no metainfo at all leaves the info button disabled — which is the honest
    # signal. Such a layer gets no file, and starts getting one the moment
    # somebody writes its description.
    if not wat_ziet_u.strip():
        skipped += 1
        continue

    def render(template, basis):
        """Fill one template with this layer's values, rooted at `basis`."""
        bronnen_html = maak_bronnen_html(
            layer_id,
            df_kaartlagen_bronnen,
            bronnen_dict,
            basis,
        )

        html = template
        html = vervang_template(html, "TITEL", escape_html(titel))
        html = vervang_template(html, "THEMA_SUBTHEMA", escape_html(thema))
        html = vervang_template(html, "WAT_ZIET_U", format_text(wat_ziet_u))
        html = vervang_template(html, "KAARTLAAGTYPE_CONTENT", format_text(berekenwijze))
        html = vervang_template(html, "AANNAMES_EN_ONZEKERHEDEN", format_text(aannames))
        html = vervang_template(html, "BRONNEN", bronnen_html)
        return html

    (OUTPUT_DIR / f"{safe_filename}.html").write_text(
        render(template_str, PREVIEW_IMG_BASE),
        encoding="utf-8",
    )

    # The fragment's inline images come from the spreadsheet's own HTML, so they
    # are rewritten after rendering rather than by the template.
    fragment = herschrijf_afbeeldingen(
        render(fragment_template_str, FRAGMENT_IMG_BASE),
        FRAGMENT_IMG_BASE,
    )
    (FRAGMENT_DIR / f"{safe_filename}.html").write_text(fragment, encoding="utf-8")
    count += 1


def kopieer_afbeeldingen():
    """
    Publish the images the fragments point at.

    Icons are a few KB each and are copied as they are. The kaartlaagtype
    diagrams are ~1300px PNGs totalling several MB, displayed a few hundred
    pixels wide — they are downscaled and converted to WebP, which takes about
    95% off with nothing visible lost.
    """
    icon_src = AFBEELDINGEN_DIR / "iconen"
    if icon_src.is_dir():
        shutil.copytree(icon_src, PUBLIC_IMG_DIR / "iconen", dirs_exist_ok=True)

    diagram_src = AFBEELDINGEN_DIR / "kaartlaagtypen"
    if not diagram_src.is_dir():
        return

    try:
        from PIL import Image
    except ModuleNotFoundError:
        sys.exit(
            "Pillow is required to convert the diagram images.\n"
            "Install it with:  pip install openpyxl pillow"
        )

    diagram_dst = PUBLIC_IMG_DIR / "kaartlaagtypen"
    diagram_dst.mkdir(parents=True, exist_ok=True)
    before = after = 0
    for png in sorted(diagram_src.glob("*.png")):
        before += png.stat().st_size
        image = Image.open(png)
        if image.width > DIAGRAM_MAX_WIDTH:
            height = round(image.height * DIAGRAM_MAX_WIDTH / image.width)
            image = image.resize((DIAGRAM_MAX_WIDTH, height), Image.LANCZOS)
        target = diagram_dst / f"{png.stem}.webp"
        image.convert("RGB").save(target, "WEBP", quality=DIAGRAM_QUALITY, method=6)
        after += target.stat().st_size

    if before:
        print(f"Afbeeldingen: {before / 1024:.0f} KB -> {after / 1024:.0f} KB")


def publiceer_fragmenten():
    """Copy the fragments to where the app serves them."""
    PUBLIC_META_DIR.mkdir(parents=True, exist_ok=True)
    for fragment in sorted(FRAGMENT_DIR.glob("*.html")):
        shutil.copy2(fragment, PUBLIC_META_DIR / fragment.name)


kopieer_afbeeldingen()
publiceer_fragmenten()

print(f"Klaar: {count} kaartlagen ({skipped} zonder beschrijving overgeslagen).")
print(f"  preview   {OUTPUT_DIR}")
print(f"  fragment  {PUBLIC_META_DIR}")
