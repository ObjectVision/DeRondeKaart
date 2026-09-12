import re
import warnings
from pathlib import Path
import openpyxl

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

MAAK_DIR = Path(__file__).resolve().parent
BASE_DIR = MAAK_DIR.parent

EXCEL_PATH = BASE_DIR / "metadata.xlsx"
TEMPLATE_PATH = BASE_DIR / "template" / "opmaak.html"
OUTPUT_DIR = BASE_DIR / "resultaten"
AFBEELDINGEN_DIR = BASE_DIR / "afbeeldingen"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def lees_sheet_data(wb, sheet_name):
    real_sheet_name = next(
        (name for name in wb.sheetnames
         if name.strip().lower() == sheet_name.lower()),
        None,
    )
    if not real_sheet_name:
        return []

    rows = list(wb[real_sheet_name].iter_rows(values_only=True))
    if not rows:
        return []

    header_idx = None
    headers = []
    for idx, row in enumerate(rows[:10]):
        cleaned = [str(c).strip() if c is not None else "" for c in row]
        if sum(1 for c in cleaned if c) >= 2:
            headers = cleaned
            header_idx = idx
            break

    if header_idx is None:
        return []

    data = []
    for row in rows[header_idx + 1:]:
        row_dict = {}
        for h, val in zip(headers, row):
            if h:
                str_val = str(val).strip() if val is not None else ""
                row_dict[h] = "" if str_val.lower() in ("nan", "none") else str_val
        if any(row_dict.values()):
            data.append(row_dict)

    return data


def escape_html(value):
    value = str(value or "")
    return (value.replace("&", "&amp;")
                 .replace("<", "&lt;")
                 .replace(">", "&gt;")
                 .replace('"', "&quot;"))


def format_text(value):
    """
    Metadata descriptions may contain HTML such as <a> and <br>.
    All links receive the standard link styling automatically.
    """
    html = str(value or "")

    def style_link(match):
        tag = match.group(0)

        # Als de link al een class heeft, voeg de standaard classes toe
        if re.search(r'\bclass\s*=', tag, re.IGNORECASE):
            tag = re.sub(
                r'class\s*=\s*"([^"]*)"',
                lambda m: f'class="{m.group(1)} underline hover:text-blue-800"',
                tag,
                count=1,
                flags=re.IGNORECASE,
            )
        else:
            tag = tag.replace(
                "<a",
                '<a class="underline hover:text-blue-800"',
                1,
            )

        # Als er nog geen style staat, voeg de primaire kleur toe
        if not re.search(r'\bstyle\s*=', tag, re.IGNORECASE):
            tag = tag.replace(
                ">",
                ' style="color: var(--primary-color);">',
                1,
            )

        return tag

    html = re.sub(
        r"<a\b[^>]*>",
        style_link,
        html,
        flags=re.IGNORECASE,
    )

    return html


def bron_logo_url(logo):
    if not logo:
        return ""
    logo = str(logo).strip()

    # Absolute URL: use unchanged.
    if re.match(r"^https?://", logo, re.I):
        return logo

    # Local logos are stored in ../afbeeldingen relative to generated HTML.
    return "../afbeeldingen/" + logo.lstrip("/\\")


def maak_bronnen_html(kaartlaag_id, df_kaartlagen_bronnen, bronnen_dict):
    gekoppelde_bronnen = [
        kb for kb in df_kaartlagen_bronnen
        if kb.get("Naam van de laag", "").strip() == kaartlaag_id
    ]

    if not gekoppelde_bronnen:
        return """
        <div class="rounded-xl border border-gray-200 bg-gray-50 p-4 text-sm text-gray-500">
          Voor deze kaartlaag zijn geen databronnen geregistreerd.
        </div>
        """

    html = []

    for kb in gekoppelde_bronnen:
        b_naam = (kb.get("Naam van de bron", "") or "").strip()
        datum = kb.get("Datum", "")
        b_info = bronnen_dict.get(b_naam, {})

        houder = b_info.get("Bronhouder", "")
        beschrijving = b_info.get("Beschrijving", "")
        url = b_info.get("URL", "")
        logo = bron_logo_url(b_info.get("Logo", ""))

        titel = escape_html(b_naam)
        houder_html = escape_html(houder)
        datum_html = escape_html(datum)
        link_html = escape_html(url)

        if link_html:
            titel_html = f"""
                  <div class="flex items-center justify-between gap-2 mb-0.5">
                    <h4 class="font-semibold text-gray-900 text-sm">{titel}</h4>
                    <a href="{link_html}" target="_blank" rel="noopener"
                       class="text-xs font-semibold hover:underline flex items-center gap-1"
                       style="color: var(--primary-color);">
                      Bekijk bron &rsaquo;
                    </a>
                  </div>
            """
        else:
            titel_html = f"""
                  <h4 class="font-semibold text-gray-900 text-sm mb-0.5">{titel}</h4>
            """

        logo_html = ""
        if logo:
            logo_html = f"""
                    <img src="{escape_html(logo)}"
                         alt="Logo {titel}"
                         class="max-h-full max-w-full object-contain"
                         onerror="this.style.display='none';">
            """

        html.append(f"""
			<tr class="border-b border-gray-200 bg-white even:bg-gray-100">
			  <td class="align-top pr-4 py-3">
				{logo_html}
			  </td>

			  <td class="align-top p-3">
				{titel_html}

				<div class="text-xs text-gray-500 mb-1">
				  <span class="font-medium text-gray-700">Bronhouder:</span> {houder_html}
				  <span aria-hidden="true"> | </span>
				  <span class="font-medium text-gray-700">Datum:</span> {datum_html}
				</div>

				<p class="text-xs leading-relaxed text-gray-600">
				  {format_text(beschrijving)}
				</p>
			  </td>
			</tr>        """)

    return "\n".join(html)


def vervang_template(html, key, content):
    pattern_if = re.compile(
        rf"\{{%\s*if\s+{re.escape(key)}\s*%\}}(.*?)\{{%\s*endif\s*%\}}",
        re.DOTALL,
    )

    if content and str(content).strip():
        html = pattern_if.sub(r"\1", html)
        html = re.sub(
            rf"\{{\{{\s*{re.escape(key)}\s*\}}\}}",
            lambda m: str(content),
            html,
        )
    else:
        html = pattern_if.sub("", html)
        html = re.sub(
            rf"\{{\{{\s*{re.escape(key)}\s*\}}\}}",
            "",
            html,
        )

    return html


# ------------------------------------------------------------
# INLEZEN
# ------------------------------------------------------------

if not EXCEL_PATH.exists():
    raise FileNotFoundError(f"Excel-bestand niet gevonden: {EXCEL_PATH}")

if not TEMPLATE_PATH.exists():
    raise FileNotFoundError(f"Template-bestand niet gevonden: {TEMPLATE_PATH}")

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

df_kaartlaag = lees_sheet_data(wb, "kaartlaag")
df_kaartlaagtypen = lees_sheet_data(wb, "kaartlaagtypen")
df_bronnen = lees_sheet_data(wb, "bronnen")
df_kaartlagen_bronnen = lees_sheet_data(wb, "kaartlagen_bronnen")

kaartlaagtypen_dict = {
    row.get("Naam van het kaartlaagtype", ""): row.get("Beschrijving", "")
    for row in df_kaartlaagtypen
    if row.get("Naam van het kaartlaagtype")
}

bronnen_dict = {
    (row.get("Naam de bron", "") or row.get("Naam van de bron", "")).strip(): row
    for row in df_bronnen
    if (row.get("Naam de bron", "") or row.get("Naam van de bron", ""))
}

template_str = TEMPLATE_PATH.read_text(encoding="utf-8")

count = 0

for row in df_kaartlaag:
    layer_id = row.get("Naam van de laag", "").strip()

    if (
        not layer_id
        or layer_id.lower().startswith("algemene informatie")
        or layer_id.lower().startswith("naam van")
    ):
        continue

    safe_filename = re.sub(r'[\\/*?:"<>|]', "_", layer_id).strip()

    titel = row.get("Titel") or layer_id
    thema = (row.get("Thema / Subthema") or "Algemeen").replace(";", " ›")
    wat_ziet_u = row.get("Wat ziet u?", "")
    kaartlaagtype_id = row.get("Gekozen kaartlaagtype", "")
    berekenwijze = kaartlaagtypen_dict.get(kaartlaagtype_id, "")
    aannames = row.get("Aannames en Onzekerheden", "")

    bronnen_html = maak_bronnen_html(
        layer_id,
        df_kaartlagen_bronnen,
        bronnen_dict,
    )

    html = template_str

    html = vervang_template(html, "TITEL", escape_html(titel))
    html = vervang_template(html, "THEMA_SUBTHEMA", escape_html(thema))
    html = vervang_template(html, "WAT_ZIET_U", format_text(wat_ziet_u))
    html = vervang_template(html, "KAARTLAAGTYPE_CONTENT", format_text(berekenwijze))
    html = vervang_template(html, "AANNAMES_EN_ONZEKERHEDEN", format_text(aannames))

    # Dynamische bronsectie
    html = vervang_template(html, "BRONNEN", bronnen_html)

    output_path = OUTPUT_DIR / f"{safe_filename}.html"
    output_path.write_text(html, encoding="utf-8")
    count += 1

print(f"Klaar: {count} HTML-bestanden aangemaakt in {OUTPUT_DIR}")
